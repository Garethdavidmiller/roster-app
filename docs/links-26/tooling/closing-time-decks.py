# THE CLOSING TIME PRESENTATIONS (30 Oct 2026, owner: "create the manager and colleague presentations"). Closing Time
# shares Short Run's weeks in everything the decks say about them — the run (5, up to 6), the weekends (7 in 26, at most
# five weeks apart), no mixed week, 22 of 26 rest-day breaks, the leave, the 42 finishes after 23:00 and the Sunday — so
# its decks are Short Run's (Second Edition's edits, then Short Run's) with what differs changed:
#   · the closers: 16:00 Monday to Saturday (7h 55m, today 8h 40m and 9h 10m), the brief it was built to (26 Oct 2026)
#     in place of the flexible rule's 15:45 — so "all 3 flexible rules" becomes two, and the decks say why;
#   · the shift times: 13 (8 worked today, 5 new), today's 13:45 kept every weekday, and on a Saturday only the closer new;
#   · the cost the closers are paid for with: three weekday 13:30 lates run to 22:30, so 104 finishes a year at 22:00 or
#     later (today 78) and eight on after 22:00 on a weekday;
#   · rest 14h 35m, weeks 27h 00m to 41h 30m, 17 of 21 weeks on one shift time, 58 starts at 06:20, four at the Saturday
#     open, and the fit chart (36.4 · 8.7 · 29.0).
# Closing Time is not on the shortlist, so the manager "Why" slide (deck-polish.py) says what it shares with the three
# shortlisted links rather than counting itself among them. Like every deck, they compare otherwise only with TODAY's link.
# Built from the Second Nature template like every deck, so run second-nature-decks.py first. From the repository root:
#   python3 docs/links-26/tooling/second-nature-decks.py && python3 docs/links-26/tooling/closing-time-decks.py \
#     && python3 docs/links-26/tooling/deck-polish.py "Closing Time" && node docs/links-26/tooling/deck-check.mjs
# Every edit refuses if the shape does not hold the text it expects; shapes are found by their text, as in Short Run's.
import importlib.util, os, re, shutil, tempfile, zipfile
import openpyxl

def load(name):
    spec = importlib.util.spec_from_file_location(name.replace('-', '_'), os.path.join(os.path.dirname(__file__), name + '.py'))
    mod = importlib.util.module_from_spec(spec); spec.loader.exec_module(mod); return mod
snd, sed, srd = load('second-nature-decks'), load('second-edition-decks'), load('short-run-decks')
Slide, notes, GREEN, runs_of, spans, escape = snd.Slide, snd.notes, snd.GREEN, snd.runs_of, snd.spans, snd.escape
put = srd.put

OUT = 'docs/links-26/presentations'
TEMPLATE = 'docs/links-26/tooling/deck-template'   # Second Nature's decks, built by second-nature-decks.py, not shipped
NAME, CODE = 'Closing Time', 'CT-26-T2'

# What Short Run's decks say that is not true of Closing Time. Searched in the text a reader sees (slide and notes
# runs, chart values), never in raw XML, where a number can be an attribute.
SHORT_RUN_ONLY = re.compile(
    r'Second Nature|Second Edition|Short Run|SN-26|SE-26|SR-26|10 weeks apart|ten weeks apart|43h 40m|25h 30m|42h 00m|42h 30m|24h 30m'
    r'|six days in a row|6 days in a row|14h 20m|18 of 21|18 of the 21|9 of 15|9 of the 15|9 of its 15|[Nn]ine of (the|its) fifteen'
    r'|\(15, not 18\)|15 shift times|[Ss]ix new|6 new|30 to 45|30–45|30 minutes|15:45–23:55|15:00–22:30|07:00–16:00|12:00–20:00'
    r'|^94$|from 78 to 94|nothing new on a Saturday|none on a Saturday is new|^29\.1$|^12\.7$|1h 13m|5 · 5 · 4|7 · 5 · 5'
    r'|15:45 weekday closers|meets all three|all 3 flexible|three flexible ones as well|starts 15:45|starts 15:15, not 14:45')

def rename(d):
    for root, _, files in os.walk(d):
        for f in files:
            if not f.endswith('.xml'): continue
            p = os.path.join(root, f); x = open(p, encoding='utf-8').read()
            y = x
            for old, new in (('Second Nature', NAME), ('Second Edition', NAME), ('Short Run', NAME), ('SN-26-F2', CODE), ('SE-26-F1', CODE), ('SR-26-F1', CODE)):
                y = y.replace(old, new)
            if y != x: open(p, 'w', encoding='utf-8').write(y)
    xp = f'{d}/ppt/embeddings/Microsoft_Excel_Worksheet1.xlsx'   # the manager deck's chart workbook: its series is still headed "Second Nature"
    if os.path.exists(xp):
        wb = openpyxl.load_workbook(xp); ws = wb.active; hit = False
        for row in ws.iter_rows():
            for cell in row:
                if cell.value in ('Second Nature', 'Second Edition', 'Short Run'): cell.value = NAME; hit = True
        if hit: wb.save(xp)

def build(src, out, edits):
    d = tempfile.mkdtemp(); zipfile.ZipFile(src).extractall(d)
    for e in edits: e(d)
    rename(d)
    left = []
    for r, _, fs in os.walk(d):
        for f in fs:
            if not f.endswith('.xml'): continue
            x = open(os.path.join(r, f), encoding='utf-8').read()
            for t in re.findall(r'<a:t>([^<]*)</a:t>|<c:v>([^<]*)</c:v>', x):
                t = t[0] or t[1]
                if SHORT_RUN_ONLY.search(t): left.append(f'{f}: {t[:90]!r}')
    if left: raise SystemExit(f'{out}: Short Run text left —\n  ' + '\n  '.join(left))
    if os.path.exists(out): os.remove(out)
    with zipfile.ZipFile(src) as zs, zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as z:
        for n in zs.namelist(): z.write(os.path.join(d, n), n)
    shutil.rmtree(d)

KNOWN = ('06:20–13:45', '06:20–14:00', '06:20–14:20', '06:20–14:50', '07:15–15:45', '08:00–16:30', '14:00–22:30', '14:30–22:00')

def colleagues(d):
    S = lambda n: Slide(f'{d}/ppt/slides/slide{n}.xml')
    s = S(2)
    put(s, '9 of 15', '8 of 13'); put(s, '6 new ones', '5 new ones')
    put(s, '30 min', '45 min'); put(s, '45 min on a Sunday', '1h 15m on a Saturday')   # deck-polish.py makes the sub-caption "Saturdays: 1h 15m"
    put(s, 'Same hours, shorter closers, weekends off never more than five weeks apart, and never more than five days in a row on the fixed rota.',
           'Same hours, a 16:00 closer every day but Sunday, weekends off never more than five weeks apart, and at most five days in a row on the fixed rota.')
    s.save()
    notes(d, 2, 'This is the whole pitch', 'This is the whole pitch on one slide. Same 35-hour week as now. Every closing shift is shorter: the weekday and Saturday closers start at 16:00, 45 minutes later than today’s weekday 15:15 and 1h 15m later than Saturday’s 14:45, and the Sunday closer starts 45 minutes later, at 15:15. Never more than five days in a row on the fixed rota — six at most when a cover week falls badly, where today it is seven and can reach nine. Seven full weekends off in every 26 weeks, against four in 20 today, and never more than five weeks apart — though each person works about two more Saturdays a year. Every week is all earlies or all lates, and 17 of the 21 working weeks keep one shift time Monday to Friday. Eight of the thirteen shift times are ones people already work, today’s 13:45 finish among them.')

    s = S(6)   # "8h 10m" and "8h 40m" are each in two rows, so the design's cell is found from its row label
    for old, new, was in (('Weekday closer — starts 15:45, not 15:15', 'Weekday closer — starts 16:00, not 15:15', '8h 10m'),
                          ('Saturday closer — starts 15:15, not 14:45', 'Saturday closer — starts 16:00, not 14:45', '8h 40m')):
        i = srd.shape(s, old); s.text(i, old, new); s.text(i + 2, was, '7h 55m')   # label, today, design
    put(s, 'Every closing shift is 30 to 45 minutes shorter, and nothing runs over 9 hours.', 'Every closer is 45 minutes to 1h 15m shorter, and nothing runs over 9 hours.')   # "closer", not "closing shift": the longer form broke "8h | 30m" across the banner's two lines
    s.save()
    notes(d, 6, 'Lead with the closers', 'Lead with the closers: they are the shifts people dislike most, and every one of them is shorter. Weekday 16:00 to 23:55 instead of 15:15 (7h 55m, not 8h 40m); Saturday 16:00 instead of 14:45 (7h 55m, not 9h 10m); Sunday 15:15 to 23:25 instead of 14:30 (8h 10m, not 8h 55m). Today’s 13:30 weekday late now runs to 22:30 with the ticket office: nine hours, where today it is eight and a half — the half hour is how the week keeps a 13:45 finish every weekday as well as the 16:00 closer. Be upfront about the two minutes on the average: the 26-week link may have no more than 219 contracted days a year, so the same 35 hours go into slightly fewer, slightly longer shifts. Nothing runs over 9 hours.')

    s = S(7)
    put(s, '94', '104'); put(s, '60', '58'); put(s, '14h 20m', '14h 35m', GREEN)
    s.save()
    notes(d, 7, 'Be straight about this one', 'Be straight about this one. About three more a year each finish at 23:00 or later, and all of that comes from the December rules: three on duty to the close every day, where today’s weekdays have two. In Closing Time the only people who finish after 23:00 are those three closers, so no link that meets the rules on 26 weeks can have fewer. Finishes at 22:00 or later go up more, from 78 to 104 a year: the rules want five still on duty after 22:00, and on a weekday three 13:30 lates run to 22:30 with the ticket office — today’s 13:30–22:00, half an hour longer, which is how the week keeps a 13:45 finish every weekday as well as the 16:00 closer. Early 06:20 starts go down a little, from 62 to 58 a year. And because the shortest gap between two shifts is 14h 35m, nobody finishes late and starts early the next morning.')

    s = S(8)
    put(s, '15 shift times in all — 9 of them you already work', '13 shift times in all — 8 of them you already work')
    put(s, 'You work these today (9)', 'You work these today (8)')
    i = next(i for i, (a, b) in enumerate(spans(s.x)) if runs_of(s.x[a:b])[:1] == ['06:20–14:00'])
    s.paras(i, '06:20–14:00', [(0, t) for t in KNOWN])
    put(s, 'New times (6)', 'New times (5)')
    i = next(i for i, (a, b) in enumerate(spans(s.x)) if runs_of(s.x[a:b])[:1] == ['Weekdays'])
    # paragraphs 0, 4 and 6 are the day headings, the rest bullets: two weekday times, Saturday's closer, three Sunday times
    s.paras(i, 'Weekdays', [(0, 'Weekdays'), (1, '13:30–22:30 — today’s 13:30 late, to 22:30 with the ticket office'), (2, '16:00–23:55 — the closer, 45 minutes later than 15:15'),
                            (4, 'Saturday'), (5, '16:00–23:55 — the closer, 1h 15m later than 14:45'),
                            (6, 'Sunday'), (7, '15:15–23:25 — the closer, 45 minutes later than 14:30'), (8, '09:00–18:00 — a new early, 9 hours'), (9, '14:30–22:30 — the ticket-office late')])
    put(s, 'Fewer shift times than today (15, not 18) — and nothing new on a Saturday.', 'Fewer shift times than today (13, not 18) — and on a Saturday only the closer is new.')
    s.save()
    notes(d, 8, 'Second Nature was built', 'Closing Time was built from today’s shift times: eight of its thirteen are worked today, and it uses fewer different times overall than today’s eighteen. Today’s 06:20–13:45 stays on every weekday and is on a Saturday too; today’s 06:20–13:35 goes. Three of the new times are shorter versions of today’s lates: the weekday and Saturday closers start at 16:00, the Sunday closer at 15:15, and the Sunday office late finishes at 22:30, not 23:25. The weekday’s 13:30–22:30 is today’s 13:30–22:00 half an hour longer, ending with the ticket office. The Sunday opener stays today’s 07:15–15:45; the Sunday 09:00–18:00 is new.')

    s = S(9)
    put(s, '14h 20m', '14h 35m', GREEN); put(s, '18 of 21', '17 of 21', GREEN)
    s.save()
    notes(d, 9, 'A full weekend is', 'A full weekend is Saturday and the following Sunday both off. Today that is 4 weekends in a 20-week cycle; Closing Time gives 7 in 26, and they are never more than five weeks apart, against seven today. Five is as close as they can be: a full weekend cannot start in a cover week or the week before one, so only 16 of the 26 weeks can begin one. Nobody works more than five days in a row on the fixed rota — today it is seven — and six is the most a badly placed cover week can make it. Every working week is all earlies or all lates — today seven weeks in 16 mix the two — and 17 of the 21 keep the same shift time Monday to Friday. Rest days mostly come in pairs: 22 of the 26 rest-day breaks are two days or more, against 13 of 17 today, and single rest days stay at four, as today. There are no six-day weeks (today there is one).')

    s = S(13)
    put(s, 'Six new shift times to learn', 'Five new shift times to learn')
    put(s, 'Several of them are shorter versions of the lates you already work today.', 'Three are shorter versions of today’s lates; today’s 13:30 late runs half an hour longer.')
    put(s, 'A few more late finishes', 'More late finishes')
    # the sheet's own figures (still on at 22:00; finishing at 23:00 or later), so the slide's "each one is on the full proposal
    # sheet too" stays true; the per-person 104 at 22:00 or later is slide 7's, and the sheet does not print it
    put(s, '42 a year each at 23:00 or later (today 39) — the rules’ minimum.', 'Eight on after 22:00 on a weekday, today four; 42 a year each after 23:00, the minimum.')
    s.save()

    s = S(14)
    old = ['Today', 'Second Nature', 'Contracted hours a week', '35h', '35h', 'Contracted days at work a year', '219', '219', 'Average shift · longest shift', '8h 14m · 9h 10m', '8h 16m · 9h 00m', 'Closing shift: weekday · Saturday · Sunday', '8h 40m · 9h 10m · 8h 55m', '8h 10m · 8h 40m · 8h 10m', 'Full weekends off', '4 in 20', '7 in 26', 'Most days worked in a row', '7 (up to 9)', '5 (up to 6)', 'Shortest gap between two shifts', '12h 30m', '14h 20m', 'Different shift times', '18', '15 (9 you know)', 'On duty weekday · Saturday · Sunday', '11–12 · 10 · 8', '15 · 14 · 10', 'Avoidable fatigue warnings in the fixed rota', '4 (up to 5)', '0']
    i = next(i for i, (a, b) in enumerate(spans(s.x)) if runs_of(s.x[a:b])[:2] == ['Today', 'Second Nature'])   # renamed after the edits
    new = list(old); new[1] = NAME; new[13] = '7h 55m · 7h 55m · 8h 10m'; new[22] = '14h 35m'; new[25] = '13 (8 you know)'
    s.runs(i, old, new); s.save()

def managers(d):
    S = lambda n: Slide(f'{d}/ppt/slides/slide{n}.xml')
    notes(d, 2, 'The nine rules are', 'The nine rules are the December 2026 rules, as confirmed verbally on 29 September 2026: at least four at the open, three to the close, five after 22:00, 10 on a Sunday, the ticket office in fixed pairs, two on the floor at all times, handovers, 8–9 hour Sundays and no more shift times than today. Behind them sit three flexible rules that are aimed for when designing but not scored on: 14 on a Saturday, five cover weeks as evenly spaced as 26 weeks allow, and weekday closers starting at 15:45. Closing Time meets the first two. Its weekday closers start at 16:00 instead — fifteen minutes later, so shorter — which is the brief it was built to (26 October 2026): today’s link on 26 lines, keeping today’s shift times where it could and a 13:45 finish every weekday. It also stays under the ceiling of 219 contracted days a year agreed on 1 October 2026, at 218.6.')

    s = S(3); put(s, '5 · 5 · 4', '5 · 4 · 4'); put(s, '7 · 5 · 5', '8 · 5 · 5'); s.save()
    notes(d, 3, 'Today’s link meets', 'Today’s link meets 4 of the 9 December rules; Closing Time meets all 9, and two of the three flexible ones (14 on a Saturday, evenly spaced cover weeks) — its weekday closers start at 16:00, not 15:45. Someone finishing at exactly 22:00 is not counted as still on after 22:00. The floor figure leaves out the ticket office, except its second person at the quiet ends of the day. Three to the close every day is the minimum, and Closing Time rosters exactly that, which is what keeps its late finishes down. On a weekday eight are still on after 22:00, three more than the rules ask: the three 13:30 lates run to 22:30 with the ticket office. A Saturday opens with four at 06:20, the rules’ minimum and today’s figure.')

    c = f'{d}/ppt/charts/chart1.xml'; x = open(c, encoding='utf-8').read()
    for old, new in (('<c:v>29.1</c:v>', '<c:v>36.4</c:v>'), ('<c:v>12.7</c:v>', '<c:v>8.7</c:v>')):
        if x.count(old) != 1: raise SystemExit(f'chart: {old} found {x.count(old)} times')
        x = x.replace(old, new)
    open(c, 'w', encoding='utf-8').write(x)
    xp = f'{d}/ppt/embeddings/Microsoft_Excel_Worksheet1.xlsx'; wb = openpyxl.load_workbook(xp); ws = wb.active; col = None
    for row in ws.iter_rows():
        for cell in row:
            if cell.value in ('Second Nature', 'Short Run'): cell.value = NAME; col = cell.column
    if col is None: raise SystemExit('chart workbook: no design column')
    for r, v in ((2, 36.4), (3, 8.7), (4, 29.0)): ws.cell(row=r, column=col).value = v
    wb.save(xp)
    notes(d, 4, 'The figure compares', 'The figure compares, hour by hour, the share of the day’s floor staff on duty with the share of the day’s train service, each train weighted by its length. 0 would be a perfect match. It is a guide to shape, not a staffing requirement — the rules on the previous slide are the requirement. Against today’s link, Closing Time’s score is about seven-tenths on weekdays, about an eighth on a Saturday and about two-fifths on a Sunday; lower is closer. The largest single gap left on a weekday is 20:00 to 22:00 — seven on the floor while the trains are quiet, because the 13:30 lates stay until 22:30 — and the morning and evening peaks are a little short of the trains.')

    notes(d, 5, 'Working weeks average', 'Working weeks average exactly 35 hours Monday to Saturday — 44,100 minutes across the 21 working weeks — so no overtime is needed to balance the Monday-to-Saturday week. Individual weeks range from 27h 00m to 41h 30m — today’s run from 25h 10m to 43h 50m, so both the heaviest and the lightest week sit inside today’s range; only the average is the contract, as today. Five cover weeks keep about the same share of the link as today’s four in 20 (one week in 5.2, against one in 5). Sunday is not contracted, so the ten Sunday duties are worked as rest-day overtime, exactly as today’s eight are. The two extra Sunday overtime duties a week come from the Sunday staffing of ten, confirmed verbally — any link that meets the rules needs them. The 26-week link, its five cover weeks and the ceiling of 219 contracted days a year were agreed on 1 October 2026; Closing Time gives 218.6.')

    s = S(6); put(s, '14h 20m', '14h 35m', GREEN); s.save()

    s = S(7)   # deck-polish.py rewrites the subtitle and the six lines; these are the pre-polish wording, true of Closing Time
    put(s, 'No run over five days on the fixed rota, and the fewest late finishes the rules allow', 'A 16:00 closer every day but Sunday, and the fewest finishes after 23:00 the rules allow')
    put(s, 'Meets all 9 December rules, and all 3 flexible ones', 'Meets all 9 December rules, and 2 of the 3 flexible ones')
    put(s, '9 of its 15 shift times are worked today', '8 of its 13 shift times are worked today')
    s.save()
    notes(d, 7, 'This is the case in one slide', 'This is the case in one slide. Closing Time answers a later brief than the shortlist (26 October 2026): today’s link on 26 lines, keeping today’s shift times where it could, a 13:45 finish every weekday and a closer starting at 16:00. It meets every December staffing rule and two of the three flexible ones: the third asks for weekday closers at 15:45, and its brief moved them to 16:00, shorter still. The weekday and Saturday closers are 7h 55m, where today they are 8h 40m and 9h 10m, and the Sunday closer is 45 minutes shorter. Eight of its thirteen shift times are worked today. Nobody works more than five days in a row on the fixed rota, its seven full weekends off in 26 weeks are never more than five weeks apart, every working week is all earlies or all lates, and 17 of the 21 keep one shift time Monday to Friday, where today it is seven weeks in 16. No working week is over 41h 30m or under 27h 00m, and the shortest rest is 14h 35m. The honest costs: to keep a 13:45 every weekday as well as the 16:00 closer, three weekday 13:30 lates run to 22:30, so finishes at 22:00 or later go from about 78 to 104 a year each; finishes at 23:00 or later go from 39 to 42, the smallest rise the December staffing allows; and five shift times are new.')

    s = S(8)
    s.runs(next(i for i, (a, b) in enumerate(spans(s.x)) if runs_of(s.x[a:b])[:1] == ['7 weekends off, at most 5 weeks apart']),
           ['7 weekends off, at most 5 weeks apart', 'At most 5 days in a row, fixed rota (today 7)', 'One turn in 18 of 21 weeks (today 7 of 16)', 'Every closing shift 30–45 minutes shorter', '9 of the 15 shift times already worked'],
           ['7 weekends off, at most 5 weeks apart', 'At most 5 days in a row, fixed rota (today 7)', 'One turn in 17 of 21 weeks (today 7 of 16)', 'Every closing shift 45 min–1h 15m shorter', '8 of the 13 shift times already worked'])
    s.runs(next(i for i, (a, b) in enumerate(spans(s.x)) if runs_of(s.x[a:b])[:1] == ['Late finishes: 42 a year (today 39)']),
           ['Late finishes: 42 a year (today 39)', '6 new shift times to learn', 'Average shift about 2 minutes longer', 'Best 14-day leave stretch 28 days, not 30', 'Four weeks off needs 15 days’ leave, not 14'],
           ['At 22:00 or later: 104 a year (today 78)', '5 new shift times to learn', 'The 13:30 late runs to 22:30, not 22:00', 'Best 14-day leave stretch 28 days, not 30', 'Four weeks off needs 15 days’ leave, not 14'])
    s.save()
    notes(d, 8, 'Days at work stay', 'Days at work stay at 219 a year (218.6 exactly; Monday to Saturday; Sunday is overtime), under the ceiling agreed on 1 October 2026. The late finish colleagues will notice is 22:00 or later, about 26 more a year each: partly the December staffing, which wants five still on after 22:00, and partly the weekday’s three 13:30 lates, which now run to 22:30 — how the week keeps a 13:45 every weekday as well as the 16:00 closer. Finishes at 23:00 or later go up about three a year, all from the December staffing — three to the close every day, where today’s weekdays have two — so no link meeting the same rules on 26 weeks could have fewer. Full weekends off are seven in 26 and never more than five weeks apart; nobody works more than five days in a row on the fixed rota; no week mixes earlies and lates, and 17 of the 21 working weeks keep one shift time. Weekly hours run from 27h 00m to 41h 30m (today 25h 10m to 43h 50m); only the average is the contract. The average shift is about two minutes longer.')

if __name__ == '__main__':
    build(f'{TEMPLATE}/Second-Nature-for-colleagues.pptx', f'{OUT}/Closing-Time-for-colleagues.pptx', [sed.colleagues, srd.colleagues, colleagues])
    build(f'{TEMPLATE}/Second-Nature-for-managers.pptx', f'{OUT}/Closing-Time-for-managers.pptx', [sed.managers, srd.managers, managers])
    print('wrote', OUT)
