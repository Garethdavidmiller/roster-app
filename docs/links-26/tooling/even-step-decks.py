# THE EVEN STEP PRESENTATIONS (31 Oct 2026, owner: "Do your suggestions" — build Even Step as In Step was built, recommend it,
# and put it to colleagues next to In Step; "Yes, keep all three"). Even Step is In Step's weeks — the same rest days,
# weekends, cover weeks and early and late weeks — with the shift times chosen again so that nothing Monday to Saturday
# runs over 8h 30m, which was the one thing about In Step the owner expected to be hard to sell ("the length of the main
# early shift"). So its decks are In Step's (Second Edition's edits, then Short Run's, then In Step's) with what differs
# changed:
#   · the shift times: 13, 8 worked today and 5 new, where In Step has 15, 9 and 6 — every early Monday to Saturday is a
#     time worked today, today's 13:45 stays every weekday and today's 06:20–14:00 on a Saturday, and the one new weekday
#     time that is not a closer is a 12:15–20:30 late;
#   · nothing over 8h 30m Monday to Saturday, where In Step has 40 weekday earlies a year over it (8h 55m and 8h 45m) and
#     today's closers run to 8h 40m and 9h 10m; the one longer shift is the Sunday 09:00–18:00, 9 hours, which all three
#     shortlisted links share;
#   · weeks of 25h 30m to 41h 30m, not 25h 30m to 41h 05m;
#   · the fit chart (32.6 · 10.0 · 29.0, not 29.8 · 9.6 · 29.0);
#   · the 55-hour week a cover week can make, 56.8 hours, not 56.3 (deck-polish.py writes it from deck-check.mjs).
# Everything else the decks quote — the run, the weekends, the rest-day breaks, the one-turn weeks, the shortest rest
# (14h 35m), the closers (8h 10m all three), the leave, the late finishes (42 after 23:00, 94 at 22:00 or later), 60 starts
# at 06:20, the staffing, the Sunday, all three flexible rules — is In Step's figure exactly (deck-check.mjs --print), and
# stays. Its manager "Why" slide (deck-polish.py) says what all three shortlisted links give, what Even Step adds against
# the other two and what it costs, and its notes set it beside both. Like every deck, they compare otherwise only with
# TODAY's link.
# Built from the Second Nature template like every deck, so run second-nature-decks.py first. From the repository root:
#   python3 docs/links-26/tooling/second-nature-decks.py && python3 docs/links-26/tooling/even-step-decks.py \
#     && python3 docs/links-26/tooling/deck-polish.py "Even Step" && node docs/links-26/tooling/deck-check.mjs
# Every edit refuses if the shape does not hold the text it expects; shapes are found by their text, as in In Step's.
import importlib.util, json, os, re, shutil, subprocess, tempfile, zipfile
import openpyxl

def load(name):
    spec = importlib.util.spec_from_file_location(name.replace('-', '_'), os.path.join(os.path.dirname(__file__), name + '.py'))
    mod = importlib.util.module_from_spec(spec); spec.loader.exec_module(mod); return mod
snd, sed, srd, isd = load('second-nature-decks'), load('second-edition-decks'), load('short-run-decks'), load('in-step-decks')
Slide, notes, GREEN, runs_of, spans, escape = snd.Slide, snd.notes, snd.GREEN, snd.runs_of, snd.spans, snd.escape
put = srd.put

OUT = 'docs/links-26/presentations'
TEMPLATE = 'docs/links-26/tooling/deck-template'   # Second Nature's decks, built by second-nature-decks.py, not shipped
NAME, CODE = 'Even Step', 'ES-26-T5'

# What In Step's decks say that is not true of Even Step, searched in the text a reader sees (slide and notes runs, chart
# values), after the rename: its two long weekday earlies and their lengths (as a pair: 8h 55m alone is today's Sunday
# closer too), its 13:35 Saturday and its 12:00–20:00, its count of times (15, 9 known, 6 new), its heaviest week, its fit
# and its 55-hour week. The figures the two share (94, 60, 18 of 21, 22 of 26, 14h 35m, 8h 10m, 5 · 5 · 4, 7 · 5 · 5) are
# deliberately absent: they are Even Step's too.
IN_STEP_ONLY = re.compile(
    r'41h 05m|^29\.8$|^9\.6$|56\.3|8h 55m and 8h 45m|up to 8h 55m|8h 45m|06:20–15:15|07:00–15:45|06:20–13:35|12:00–20:00'
    r'|40 in all|new weekday earlies|[Ll]onger weekday earl|also a weekday time|a new late, ending with the ticket office'
    r'|^15 \(9 you know\)$|^9 of 15$|^6 new ones$|[Nn]ine of (its|the) fifteen|\(15, not 18\)|9 of the 15'
    r'|[Ss]ix new shift times|6 new shift times|^15 shift times in all|You work these today \(9\)|New times \(6\)')
# The template's sentences that name their design ("today or in Second Nature") read wrongly once that name is Even Step.
# In Step's builder rewrote each of them; this refuses any that survives.
AWKWARD = re.compile(r'\b[Ii]n Even Step\b')

def rename(d):
    for root, _, files in os.walk(d):
        for f in files:
            if not f.endswith('.xml'): continue
            p = os.path.join(root, f); x = open(p, encoding='utf-8').read()
            y = x
            for old, new in (('Second Nature', NAME), ('Second Edition', NAME), ('Short Run', NAME), ('In Step', NAME),
                             ('SN-26-F2', CODE), ('SE-26-F1', CODE), ('SR-26-F1', CODE), ('IS-26-T4', CODE)):
                y = y.replace(old, new)
            if y != x: open(p, 'w', encoding='utf-8').write(y)
    xp = f'{d}/ppt/embeddings/Microsoft_Excel_Worksheet1.xlsx'   # the manager deck's chart workbook
    if os.path.exists(xp):
        wb = openpyxl.load_workbook(xp); ws = wb.active; hit = False
        for row in ws.iter_rows():
            for cell in row:
                if cell.value in ('Second Nature', 'Second Edition', 'Short Run', 'In Step'): cell.value = NAME; hit = True
        if hit: wb.save(xp)

# Before the rename the template's own name, Second Nature, is the only design a slide or a note may name: another
# design's name at that point is a sentence ABOUT that design, carried over, which the rename would quietly turn into a
# sentence about this one (Closing Time's builder, 30 Oct 2026 review; in git). In Step is in the list too: its builder's
# own sentences name it, and each has to be rewritten here, not renamed.
OTHER_NAMES = re.compile(r'Second Edition|Short Run|Closing Time|In Step|SE-26-F1|SR-26-F1|CT-26-T2|IS-26-T4')
# The deliberate mentions of another design — the managers' notes set Even Step beside In Step and Short Run — are written
# with a no-break space, as second-nature-decks.py writes Familiar Nine, so the check above passes them and the rename
# leaves them alone. In Step's own figures quoted beside Even Step's (its 41h 05m week) carry the same mark, for
# IN_STEP_ONLY's sake.
IS = 'In\u00a0Step'
SR = 'Short\u00a0Run'

def build(src, out, edits):
    d = tempfile.mkdtemp(); zipfile.ZipFile(src).extractall(d)
    for e in edits: e(d)
    carried = [f'{f}: {t[:90]!r}' for r, _, fs in os.walk(d) for f in fs if f.endswith('.xml')
               for t in re.findall(r'<a:t>([^<]*)</a:t>', open(os.path.join(r, f), encoding='utf-8').read()) if OTHER_NAMES.search(t)]
    if carried: raise SystemExit(f'{out}: another design named before the rename —\n  ' + '\n  '.join(carried))
    rename(d)
    left = []
    for r, _, fs in os.walk(d):
        for f in fs:
            if not f.endswith('.xml'): continue
            x = open(os.path.join(r, f), encoding='utf-8').read()
            for t in re.findall(r'<a:t>([^<]*)</a:t>|<c:v>([^<]*)</c:v>', x):
                t = t[0] or t[1]
                if IN_STEP_ONLY.search(t): left.append(f'{f}: In Step text {t[:90]!r}')
                if AWKWARD.search(t): left.append(f'{f}: renamed into "in Even Step" {t[:90]!r}')
    if left: raise SystemExit(f'{out}: text that is not Even Step\'s —\n  ' + '\n  '.join(left))
    if os.path.exists(out): os.remove(out)
    with zipfile.ZipFile(src) as zs, zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as z:
        for n in zs.namelist(): z.write(os.path.join(d, n), n)
    shutil.rmtree(d)

KNOWN = ('06:20–13:45', '06:20–14:00', '06:20–14:20', '06:20–14:50', '07:15–15:45', '08:00–16:30', '14:00–22:30', '14:30–22:00')
CLOSERS = '8h 10m · 8h 10m · 8h 10m'

def colleagues(d):
    S = lambda n: Slide(f'{d}/ppt/slides/slide{n}.xml')
    s = S(2)
    put(s, '9 of 15', '8 of 13'); put(s, '6 new ones', '5 new ones')
    # the pitch line leads with what Even Step is for: no long day Monday to Saturday (the weekends are the tile beside it)
    put(s, 'Same hours, shorter closers, weekends off never more than five weeks apart, and never more than five days in a row on the fixed rota.',
           'Same hours, shorter closers, nothing over 8h\u00a030m Monday to Saturday, and never more than five days in a row on the fixed rota.')   # no-break: a duration never splits
    s.save()
    notes(d, 2, 'This is the whole pitch', 'This is the whole pitch on one slide. Same 35-hour week as now. Nothing Monday to Saturday is longer than 8h 30m — today the closers run to 8h 40m and 9h 10m — and every early shift Monday to Saturday is a time people already work, though two of the weekday ones are worked today only at the weekend. Never more than five days in a row on the fixed rota — seven at most when a cover week falls badly and its Sunday is worked as overtime, where today it is seven and can reach nine. Seven full weekends off in every 26 weeks, against four in 20 today, and never more than five weeks apart — though each person works about two more Saturdays a year. Every working week is all earlies or all lates, and 18 of the 21 working weeks keep one shift time Monday to Friday. Eight of the thirteen shift times are ones people already work, today’s 13:45 finish among them. And every closing shift is shorter: the weekday closer by 30 minutes, the Saturday closer by an hour and the Sunday closer by 45 minutes.')

    notes(d, 4, 'Monday to Saturday, counting', f'Monday to Saturday, counting a cover week as four days worked, the same way the clerks do. Sundays are overtime by agreement, so they are not in these figures, today’s or {NAME}’s. Exactly, {NAME} works out at 218.6 days a year: just under today’s 219, which is also the most the 26-week link is allowed (set on 1 October 2026; a year here is 365 days, as the ceiling’s is). The average shift is about two minutes longer, which is on slide 6.')

    notes(d, 6, 'Lead with the closers', 'Lead with the closers: they are the shifts people dislike most, and every one of them is shorter. Weekday 15:45 to 23:55 instead of 15:15 (8h 10m, not 8h 40m); Saturday 15:45 too, instead of 14:45 (8h 10m, not 9h 10m); Sunday 15:15 to 23:25 instead of 14:30 (8h 10m, not 8h 55m). Nothing Monday to Saturday runs over 8h 30m: every early from Monday to Saturday is a time worked today, and the weekday lates are 8h 30m and 8h 15m. But the time the closers lose goes onto the other shifts, so more of them run to the full 8h 30m than today — four of the seven weekday earlies, where today two of six do. The one shift over 8h 30m anywhere is the Sunday 09:00–18:00, 9 hours, worked as overtime about twice a year each. Be upfront about the two minutes on the average: the 26-week link may have no more than 219 contracted days a year, so the same 35 hours go into slightly fewer, slightly longer shifts. Nothing runs over 9 hours.')

    notes(d, 7, 'Be straight about this one', f'Be straight about this one. About three more a year each finish at 23:00 or later, and all of that comes from the December rules: three on duty to the close every day, where today has two Monday to Saturday. With {NAME} the only people who finish after 23:00 are those three closers, so no link that meets the rules on 26 weeks can have fewer. Finishes at 22:00 or later go up more, from 78 to 94 a year: the rules also want five still on duty after 22:00, and two weekday lates finish at 22:30 with the ticket office, on the office’s own 14:00–22:30. Early 06:20 starts hardly change. And because the shortest gap between two shifts is 14h 35m, nobody finishes late and starts early the next morning.')

    s = S(8)
    put(s, '15 shift times in all — 9 of them you already work', '13 shift times in all — 8 of them you already work')
    put(s, 'You work these today (9)', 'You work these today (8)'); put(s, 'New times (6)', 'New times (5)')
    i = next(i for i, (a, b) in enumerate(spans(s.x)) if runs_of(s.x[a:b])[:1] == ['06:20–13:35'])
    s.paras(i, '06:20–13:35', [(0, t) for t in KNOWN])
    i = next(i for i, (a, b) in enumerate(spans(s.x)) if runs_of(s.x[a:b])[:1] == ['Weekdays'])
    # paragraphs 0, 4 and 6 are the day headings, the rest bullets (In Step's ten); nine here, the weekday having two
    s.paras(i, 'Weekdays', [(0, 'Weekdays'), (1, '12:15–20:30 — a new late, 8h 15m'), (2, '15:45–23:55 — the closer, 30 minutes later than 15:15'),
                            (4, 'Saturday'), (5, '15:45–23:55 — the weekday closer, 1 hour later than 14:45'),
                            (6, 'Sunday'), (7, '15:15–23:25 — the closer, 45 minutes later than 14:30'), (8, '09:00–18:00 — a new early, 9 hours'), (9, '14:30–22:30 — the ticket-office late')])
    put(s, 'Fewer shift times than today (15, not 18) — and on a Saturday only the closer is new.', 'Fewer shift times than today (13, not 18) — and no early Monday to Saturday is new.')
    s.save()
    notes(d, 8, 'In Step was built', f'{NAME} was built keeping the times people already work wherever that cost little, and keeping every early Monday to Saturday one of them. Eight of its thirteen are worked today, and it uses fewer different times overall than today’s eighteen. Today’s 06:20–13:45 stays on every weekday, and today’s 06:20–14:00 on a Saturday. Three of the new times are shorter versions of today’s lates: the weekday and Saturday closers start at 15:45 and the Sunday closer at 15:15, and the Sunday office late finishes at 22:30, not 23:25. The weekday 12:15–20:30 is a new late, 8h 15m, and the Sunday 09:00–18:00 is new. Some known times move day: the weekday earlies 06:20–14:50 and 07:15–15:45 are worked today only on a Saturday and a Sunday, and the Saturday’s 06:20–14:20, 08:00–16:30 and 14:00–22:30 only on weekdays. The full list is on the slide. The Sunday opener stays today’s 07:15–15:45.')

    notes(d, 9, 'A full weekend is', f'A full weekend is Saturday and the following Sunday both off. Today that is 4 weekends in a 20-week cycle; {NAME} gives 7 in 26, and they are never more than five weeks apart, against seven today. Five is as close as they can be: a full weekend cannot start in a cover week or the week before one, so only 16 of the 26 weeks can begin one. Nobody works more than five days in a row on the fixed rota — today it is seven — and seven is the most a badly placed cover week can make it, with its Sunday worked as overtime. Every working week is all earlies or all lates — today seven weeks in 16 mix the two — and 18 of the 21 keep the same shift time Monday to Friday. Rest days mostly come in pairs: 22 of the 26 rest-day breaks are two days or more, against 13 of 17 today, and single rest days stay at four, as today. There are no six-day weeks (today there is one).')

    s = S(13)
    put(s, 'Six new shift times to learn', 'Five new shift times to learn')
    put(s, 'Three are shorter versions of today’s lates; two\u00a0are longer weekday earlies.', 'Three are shorter versions of today’s lates; no\u00a0early Monday to Saturday is new.')   # no-break: no lone "no" at a line end
    s.save()

    s = S(14)
    old = ['Today', 'In Step', 'Contracted hours a week', '35h', '35h', 'Contracted days at work a year', '219', '219', 'Average shift · longest shift', '8h 14m · 9h 10m', '8h 16m · 9h 00m', 'Closing shift: weekday · Saturday · Sunday', '8h 40m · 9h 10m · 8h 55m', CLOSERS, 'Full weekends off', '4 in 20', '7 in 26', 'Most days worked in a row', '7 (up to 9)', '5 (up to 7)', 'Shortest gap between two shifts', '12h 30m', '14h 35m', 'Different shift times', '18', '15 (9 you know)', 'On duty weekday · Saturday · Sunday', '11–12 · 10 · 8', '15 · 14 · 10', 'Avoidable fatigue warnings in the fixed rota', '4 (up to 5)', '0']
    i = next(i for i, (a, b) in enumerate(spans(s.x)) if runs_of(s.x[a:b])[:2] == ['Today', 'In Step'])   # In Step's header, typed by its builder
    new = list(old); new[1] = NAME; new[25] = '13 (8 you know)'
    s.runs(i, old, new); s.save()

def managers(d):
    S = lambda n: Slide(f'{d}/ppt/slides/slide{n}.xml')
    c = f'{d}/ppt/charts/chart1.xml'; x = open(c, encoding='utf-8').read()
    for old, new in (('<c:v>29.8</c:v>', '<c:v>32.6</c:v>'), ('<c:v>9.6</c:v>', '<c:v>10</c:v>')):   # the labels' own format writes 10 as 10.0
        if x.count(old) != 1: raise SystemExit(f'chart: {old} found {x.count(old)} times')
        x = x.replace(old, new)
    open(c, 'w', encoding='utf-8').write(x)
    xp = f'{d}/ppt/embeddings/Microsoft_Excel_Worksheet1.xlsx'; wb = openpyxl.load_workbook(xp); ws = wb.active; col = None
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == 'In Step': cell.value = NAME; col = cell.column   # In Step's builder headed the series with its name
    if col is None: raise SystemExit('chart workbook: no design column')
    for r, v in ((2, 32.6), (3, 10.0), (4, 29.0)): ws.cell(row=r, column=col).value = v
    wb.save(xp)
    notes(d, 4, 'The figure compares', f'The figure compares, hour by hour, the share of the day’s floor staff on duty with the share of the day’s train service, each train weighted by its length. 0 would be a perfect match. It is a guide to shape, not a staffing requirement — the rules on the previous slide are the requirement. Against today’s link, {NAME}’s score is a little over three-fifths on weekdays, about a seventh on a Saturday and about two-fifths on a Sunday; lower is closer. Its weekday is the loosest of the three shortlisted links ({IS}’s 29.8, {SR}’s 29.1): holding every shift Monday to Saturday to 8h 30m or less costs a little of the weekday shape.')

    notes(d, 5, 'Working weeks average', f'Working weeks average exactly 35 hours Monday to Saturday — 44,100 minutes across the 21 working weeks — so no overtime is needed to balance the Monday-to-Saturday week. Individual weeks range from 25h 30m to 41h 30m — today’s run from 25h 10m to 43h 50m, so both the heaviest and the lightest week sit inside today’s range; only the average is the contract, as today. Five cover weeks keep about the same share of the link as today’s four in 20 (one week in 5.2, against one in 5). Sunday is not contracted, so the ten Sunday duties are worked as rest-day overtime, exactly as today’s eight are. The two extra Sunday overtime duties a week come from the Sunday staffing of ten, confirmed verbally — any link that meets the rules needs them. The 26-week link, its five cover weeks and the ceiling of 219 contracted days a year were set on 1 October 2026; {NAME} gives 218.6.')

    # slide 7's subtitle and six lines are deck-polish.py's; the notes set Even Step beside the other two shortlisted links
    notes(d, 7, 'This is the case in one slide', f'This is the case in one slide. {NAME} (31 October 2026) is {IS}’s weeks — the same rest days, weekends, cover weeks and early and late weeks — with the shift times chosen again so that nothing Monday to Saturday runs over 8h 30m. It is the recommended design, on the shortlist with {IS} and {SR}, and this slide sets it beside both. It meets every December staffing rule and all three flexible ones. Eight of its thirteen shift times are worked today, every early Monday to Saturday among them (two of the weekday earlies today only at the weekend), and today’s 13:45 finish stays on every weekday. Nobody works more than five days in a row on the fixed rota, where today it is seven; a cover week placed badly, with its Sunday worked as overtime, can make it seven, never more. Every working week is all earlies or all lates, and 18 of the 21 keep one shift time Monday to Friday; today seven weeks in 16 do both. Its seven full weekends off in 26 weeks are never more than five weeks apart, which is as close as the five cover weeks allow; today it is four in 20, up to seven weeks apart. Its working weeks run from 25h 30m to 41h 30m (today 25h 10m to 43h 50m), and the shortest rest is 14h 35m. Against the other two it has the same run, weekends, rest-day breaks, weeks on one shift time, late finishes and leave. What it adds: no shift over 8h 30m Monday to Saturday, where {IS} has 40 a year each, {SR} 26 and today 31 (today’s closers), and the fewest shift times, 13 against their 15. What it gives up: the weekday lates, the longest of the three at 8h 27m on average (closers aside), near today’s 8h 30m, where {IS}’s are 8h 12m and {SR}’s 8h 00m — the time {IS} puts on its earlies goes on {NAME}’s lates instead; the working day moves further from one week to the next (1h 31m on average, against 1h 13m for both; today 4h 00m); its weekday fit to the trains is the loosest of the three (32.6, against {IS}’s 29.8 and {SR}’s 29.1; today 52.0); and its heaviest week, 41h 30m, is 25 minutes heavier than {IS}’s, though an hour lighter than {SR}’s. The honest costs against today: finishes at 23:00 or later go from about 39 to 42 a year each, the smallest rise possible — the December staffing needs three to the close every day, and only those three finish after 23:00 — and finishes at 22:00 or later from about 78 to 94; each person works about two more Saturdays a year; and five shift times are new. Every closing shift is 30 to 60 minutes shorter than today’s.')

    s = S(8)
    # the welcome list leads with what Even Step is for and gives up the one-turn weeks (on the colleagues' slide 9 still);
    # the worry In Step had in its third place, its longer weekday earlies, is not Even Step's, so the two more Saturdays
    # (the colleagues' slide 13) take it
    s.runs(next(i for i, (a, b) in enumerate(spans(s.x)) if runs_of(s.x[a:b])[:1] == ['7 weekends off, at most 5 weeks apart']),
           ['7 weekends off, at most 5 weeks apart', 'At most 5 days in a row, fixed rota (today 7)', 'One turn in 18 of 21 weeks (today 7 of 16)', 'Every closing shift 30–60 minutes shorter', '9 of the 15 shift times already worked'],
           ['7 weekends off, at most 5 weeks apart', 'At most 5 days in a row, fixed rota (today 7)', 'Longest shift Mon–Sat 8h 30m, not 9h 10m', 'Every closing shift 30–60 minutes shorter', '8 of the 13 shift times already worked'])   # one line each
    s.runs(next(i for i, (a, b) in enumerate(spans(s.x)) if runs_of(s.x[a:b])[:1] == ['Late finishes: 42 a year (today 39)']),
           ['Late finishes: 42 a year (today 39)', '6 new shift times to learn', 'Longer weekday earlies, up to 8h 55m', 'Best 14-day leave stretch 28 days, not 30', 'Four weeks off needs 15 days’ leave, not 14'],
           ['Late finishes: 42 a year (today 39)', '5 new shift times to learn', 'Two more Saturdays a year (28, today 26)', 'Best 14-day leave stretch 28 days, not 30', 'Four weeks off needs 15 days’ leave, not 14'])
    s.save()
    notes(d, 8, 'Days at work stay', 'Days at work stay at 219 a year (218.6 exactly; Monday to Saturday; Sunday is overtime), under the ceiling set on 1 October 2026. Late finishes average about three more a year each, all of them from the December staffing — three to the close every day, where today has two Monday to Saturday — so no link meeting the same rules on 26 weeks could have fewer. Full weekends off are seven in 26 and never more than five weeks apart; nobody works more than five days in a row on the fixed rota; no week mixes earlies and lates, and 18 of the 21 working weeks keep one shift time. Weekly hours run from 25h 30m to 41h 30m (today 25h 10m to 43h 50m); only the average is the contract. No shift Monday to Saturday is over 8h 30m, where today the closers run to 8h 40m and 9h 10m; the one longer shift is the Sunday 09:00–18:00, 9 hours. The average shift is about two minutes longer, and each person works about two more Saturdays a year.')

def check_typed():
    """THE TYPED FIGURES, CHECKED, as In Step's builder checks its own: the slides' shift-time lists, labels and notes are typed
    in this file, and deck-check.mjs reads only table rows, so each is compared here with the grids and with deck-check.mjs's
    figures before a deck is built."""
    tool = os.path.dirname(os.path.abspath(__file__))
    fig = json.loads(subprocess.check_output(['node', os.path.join(tool, 'deck-check.mjs'), '--print'], text=True)); F, T, I, S = fig[NAME], fig['today'], fig['In Step'], fig['Short Run']
    today = json.loads(subprocess.check_output(['node', '-e', "import('./report-data.mjs').then(m => process.stdout.write(JSON.stringify(m.today().patterns)))"], cwd=tool, text=True))
    g = json.load(open(os.path.join(tool, '..', 'proposals', f'{NAME.replace(" ", "-")}-{CODE}.json'), encoding='utf-8')); g = g.get('patterns', g)
    grid = lambda f: (lambda j: j.get('patterns', j))(json.load(open(os.path.join(tool, f), encoding='utf-8')))
    ins, sr = grid('in-step.json'), grid('short-run.json')
    timed = lambda v: bool(re.fullmatch(r'\d\d:\d\d-\d\d:\d\d', v or ''))
    mins = lambda v: (int(v[6:8]) * 60 + int(v[9:]) - int(v[:2]) * 60 - int(v[3:5])) % 1440
    WK, MS, ALLD = ('mon', 'tue', 'wed', 'thu', 'fri'), ('mon', 'tue', 'wed', 'thu', 'fri', 'sat'), ('sun', 'mon', 'tue', 'wed', 'thu', 'fri', 'sat')
    times = lambda p, days: {r[d] for r in p.values() for d in days if timed(r[d])}
    starts = lambda p, days, end: sorted({r[d][:5] for r in p.values() for d in days if timed(r[d]) and r[d][6:] == end})
    per_day = lambda p, t, days: {sum(1 for r in p.values() if r[d] == t) for d in days}   # how many work time t, day by day
    a_year = lambda p, test: round(sum(1 for r in p.values() for d in ALLD if timed(r[d]) and test(r[d], d)) * 52 / len(p))
    dash = lambda v: sorted(t.replace('-', '–') for t in v)
    gT, toT = times(g, ALLD), times(today, ALLD); known, new = gT & toT, gT - toT
    fit = lambda r: f'{r[0]} · {r[1]} · {r[2]}'
    frac = lambda a, b: float(a) / float(b)
    hm = lambda m: f'{m // 60}h {m % 60:02d}m'
    H = lambda v: int(v.split('h')[0]) * 60 + int(v.split('h')[1].strip().rstrip('m'))
    src = open(os.path.join(tool, 'fresh.mjs'), encoding='utf-8').read()
    SL = re.findall(r"'([^']+)'", re.search(r'export const SHORTLIST = \[([^\]]*)\]', src).group(1))
    REC = re.search(r"export const RECOMMENDED = '([^']+)'", src).group(1)
    long_ms = lambda p: a_year(p, lambda s, d: d != 'sun' and mins(s) > 510)                       # any shift over 8h 30m, Monday to Saturday
    late_avg = lambda p: (lambda l: round(sum(l) / len(l)))([mins(r[d]) for r in p.values() for d in WK if timed(r[d]) and int(r[d][:2]) >= 11 and r[d][6:] != '23:55'])   # the sheet's lateAvg
    per_wk_earlies = lambda p: {(sum(1 for r in p.values() if timed(r[d]) and early(r[d])), sum(1 for r in p.values() if timed(r[d]) and early(r[d]) and mins(r[d]) == 510)) for d in WK}
    times_on = lambda p, t: {d for r in p.values() for d in ALLD if r[d] == t}   # the days a time is worked
    early = lambda s: int(s[:2]) < 11
    shape = lambda p: [[('C' if r[d] == 'SPARE' else 'R' if not timed(r[d]) else 'E' if early(r[d]) else 'L') for d in ALLD] for k, r in sorted(p.items(), key=lambda kv: int(kv[0]))]
    checks = [   # (what the slides say, what the data says, what is typed)
        ('Even Step is In Step’s weeks: the same rest days, cover weeks and early and late days, line by line', shape(g) == shape(ins), True),
        ('the times worked today, slide 8', dash(known), sorted(KNOWN)),
        ('13 shift times, 8 worked today, 5 new; today 18, slides 2, 8, 13 and 14', [len(gT), len(known), len(new), len(toT)], [13, 8, 5, 18]),
        ('the new weekday times, slide 8', dash(new & times(g, WK)), ['12:15–20:30', '15:45–23:55']),
        ('the new Saturday time, slide 8', dash(new & times(g, ('sat',))), ['15:45–23:55']),
        ('the new Sunday times, slide 8', dash(new & times(g, ('sun',))), ['09:00–18:00', '14:30–22:30', '15:15–23:25']),
        ('no early Monday to Saturday is new, slides 8 and 13 and the notes', sorted(t for t in new if early(t) and t in times(g, MS)), []),
        ('the new late 8h 15m and the Sunday 9 hours, slide 8', [mins('12:15-20:30'), mins('09:00-18:00')], [495, 540]),
        ('nothing Monday to Saturday over 8h 30m; the weekday lates 8h 30m, 8h 15m and the 8h 10m closer, colleagues’ notes 6',
         [max(mins(t) for t in times(g, MS)), sorted({mins(t) for t in times(g, WK) if not early(t)})], [510, [490, 495, 510]]),
        ('the one shift over 8h 30m anywhere is the Sunday 09:00–18:00, about twice a year each, notes 6 and 8',
         [sorted(t for t in gT if mins(t) > 510), per_day(g, '09:00-18:00', ('sun',)), a_year(g, lambda s, d: mins(s) > 510)], [['09:00-18:00'], {1}, 2]),
        ('today’s longest Monday to Saturday 9h 10m, its closers 8h 40m and 9h 10m the only ones over 8h 30m, 31 a year each, slide 2’s notes, managers’ 7 and 8',
         [hm(max(mins(t) for t in times(today, MS))), sorted(mins(t) for t in times(today, MS) if mins(t) > 510), long_ms(today)], ['9h 10m', [520, 550], 31]),
        ('In Step 40 a year over 8h 30m Monday to Saturday, Short Run 26, Even Step none, managers’ notes 7', [long_ms(ins), long_ms(sr), long_ms(g)], [40, 26, 0]),
        ('closers from 15:45 Monday to Saturday, today 15:15 and 14:45, slides 2, 6 and 8', [starts(g, WK, '23:55'), starts(g, ('sat',), '23:55'), starts(today, WK, '23:55'), starts(today, ('sat',), '23:55')],
         [['15:45'], ['15:45'], ['15:15'], ['14:45']]),
        ('the Sunday closer from 15:15, today 14:30, slides 6 and 8', [starts(g, ('sun',), '23:25'), starts(today, ('sun',), '23:25')], [['15:15'], ['14:30']]),
        ('today’s 13:45 on every weekday and its 06:20–14:00 on a Saturday, notes 2 and 8',
         [all(n >= 1 for n in per_day(g, '06:20-13:45', WK)), per_day(g, '06:20-14:00', ('sat',)), min(per_day(today, '06:20-14:00', ('sat',))) >= 1], [True, {1}, True]),
        ('four weekday 14:00–22:30s: the office pair and two lates finishing with it, notes 7', [per_day(g, '14:00-22:30', WK), per_day(g, '14:30-22:30', WK)], [{4}, {0}]),
        ('the Sunday office late is 14:30–22:30, a pair, slide 8', per_day(g, '14:30-22:30', ('sun',)), {2}),
        ('the figures on slides 2, 4, 6, 7, 9, 14 and the manager slides',
         [F['rest'], F['closers'], T['closers'], F['heavy'], F['light'], T['heavy'], T['light'], F['late22'], T['late22'], F['late23'], T['late23'], F['open0620'], F['times'], F['oneTurn'], F['at22'], F['open'], F['run'], F['sat'], T['sat']],
         ['14h 35m', CLOSERS, '8h 40m · 9h 10m · 8h 55m', '41h 30m', '25h 30m', '43h 50m', '25h 10m', '94', '78', '42', '39', '60', '13 (8 you know)', '18 of 21', '7 · 5 · 5', '5 · 5 · 4', '5 (up to 7)', '28', '26']),
        ('the fit chart, manager slide 4', fit([F['fitWk'], F['fitSat'], F['fitSun']]), '32.6 · 10.0 · 29.0'),
        # the notes word the fit against today's as a fraction: "a little over three-fifths" (0.6–0.65), "about a seventh"
        # and "about two-fifths" (each within a hundredth of 1/7 and 2/5)
        ('the fit as fractions of today’s, manager notes 4',
         [0.6 < frac(F['fitWk'], T['fitWk']) < 0.65, abs(frac(F['fitSat'], T['fitSat']) - 1 / 7) < 0.01, abs(frac(F['fitSun'], T['fitSun']) - 2 / 5) < 0.01], [True, True, True]),
        ('the weekday fit the loosest of the three: In Step 29.8, Short Run 29.1, today 52.0, managers’ notes 4 and 7',
         [I['fitWk'], S['fitWk'], T['fitWk'], float(F['fitWk']) > max(float(I['fitWk']), float(S['fitWk']))], ['29.8', '29.1', '52.0', True]),
        ('against the other two: the same run, weekends, rest-day breaks, weeks on one shift time, late finishes and leave, managers’ notes 7',
         [[F[k] == o[k] for o in (I, S) for k in ('run', 'weekends', 'weekendGap', 'breaks', 'oneTurn', 'late23', 'late22', 'leaveBest', 'leaveAvg', 'leaveWorst', 'leaveFour')]], [[True] * 22]),
        ('the fewest shift times, 13 against their 15, managers’ notes 7', [I['times'], S['times']], ['15 (9 you know)', '15 (9 you know)']),
        ('the step 1h 31m, both the others 1h 13m, today 4h 00m, managers’ notes 7', [F['step'], I['step'], S['step'], T['step']], ['1h 31m', '1h 13m', '1h 13m', '4h 00m']),
        ('the heaviest week 25 minutes over In Step’s and an hour under Short Run’s, managers’ notes 7', [H(F['heavy']) - H(I['heavy']), H(S['heavy']) - H(F['heavy'])], [25, 60]),
        ('the weekday lates on average, closers aside: Even Step 8h 27m, In Step 8h 12m, Short Run 8h 00m, today 8h 30m, managers’ notes 7', [late_avg(g), late_avg(ins), late_avg(sr), late_avg(today)], [507, 492, 480, 510]),
        ('four of the seven weekday earlies 8h 30m, today two of six, every weekday, colleagues’ notes 6', [per_wk_earlies(g), per_wk_earlies(today)], [{(7, 4)}, {(6, 2)}]),
        ('the known times that move day: weekday earlies today only at the weekend, Saturday times today only on weekdays, notes 2, 8 and the managers’ 7',
         [sorted(t for t in times(g, WK) if early(t) and t in toT and not times_on(today, t) & set(WK)), sorted(t for t in times(g, ('sat',)) if t in toT and times_on(today, t) <= set(WK))],
         [['06:20-14:50', '07:15-15:45'], ['06:20-14:20', '08:00-16:30', '14:00-22:30']]),
        ('Even Step recommended, on the shortlist with In Step and Short Run', [REC, sorted(SL)], [NAME, ['Even Step', 'In Step', 'Short Run']]),
    ]
    bad = [f'{what}: the data says {got}, the slides say {typed}' for what, got, typed in checks if got != typed]
    if bad: raise SystemExit('even-step-decks: a typed figure no longer matches the rota —\n  ' + '\n  '.join(bad))

if __name__ == '__main__':
    check_typed()
    build(f'{TEMPLATE}/Second-Nature-for-colleagues.pptx', f'{OUT}/Even-Step-for-colleagues.pptx', [sed.colleagues, srd.colleagues, isd.colleagues, colleagues])
    build(f'{TEMPLATE}/Second-Nature-for-managers.pptx', f'{OUT}/Even-Step-for-managers.pptx', [sed.managers, srd.managers, isd.managers, managers])
    print('wrote', OUT)
