# THE SECOND NATURE PRESENTATIONS (2 Oct 2026, owner: "use the Familiar Nine and Right Away presentations as a basis").
# The two Familiar Nine decks (../../links-24/presentations/, the template Right Away's were made from too) are copied
# and every Familiar-Nine-specific word and figure is replaced, shape by shape, so the design is theirs unchanged.
# The figures were not typed from memory: they are deck-check.mjs --print (the sheets' own counts) and leave.mjs (the
# leave model, checked against the 24-line decks' published figures), and deck-check.mjs reads the finished decks back
# against the sheets. Run from the repository root:
#   python3 docs/links-26/tooling/second-nature-decks.py && node docs/links-26/tooling/deck-check.mjs
# A shape is addressed by its index in the slide (every <p:sp>/<p:graphicFrame>, in document order), and an edit
# refuses to apply when the shape does not hold the text it expects, so a changed template fails loudly.
import re, shutil, zipfile, tempfile, os, io
from xml.sax.saxutils import escape
import openpyxl

SRC = 'docs/links-24/presentations'
# Second Nature's decks are the TEMPLATE every shipped deck is built from, and are no longer shipped themselves (owner,
# 3 Oct 2026: only the shortlist's presentations are kept). They are built into a gitignored folder and read from there.
OUT = 'docs/links-26/tooling/deck-template'
NAME, CODE = 'Second Nature', 'SN-26-F2'
NAVY, GREEN, AMBER = '001E3C', '1E7B4B', 'A15C00'
# A deliberate mention of Familiar Nine (Second Nature is compared with it) is written with a no-break space, so the
# rename below leaves it alone and the leftover check below does not mistake it for template text.
F9 = 'Familiar\u00a0Nine'

SHAPE = re.compile(r'<p:(sp|graphicFrame)>.*?</p:(sp|graphicFrame)>', re.S)
AT = re.compile(r'<a:t>([^<]*)</a:t>')

def spans(x): return [m.span() for m in SHAPE.finditer(x)]
def runs_of(sp): return [m.group(1) for m in AT.finditer(sp)]

class Slide:
    def __init__(self, path): self.path = path; self.x = open(path, encoding='utf-8').read()
    def _shape(self, i): a, b = spans(self.x)[i]; return a, b, self.x[a:b]
    def _put(self, i, sp): a, b, _ = self._shape(i); self.x = self.x[:a] + sp + self.x[b:]
    def text(self, i, old, new, color=None):
        """Replace a one-run shape's text, checking it held `old`; optionally recolour it. A newline in `new` is a
        DELIBERATE line break (<a:br/>): the tile captions use it so a two-word line never becomes a lone last word."""
        a, b, sp = self._shape(i); r = runs_of(sp)
        if [escape(old)] != r: raise SystemExit(f'{self.path} shape {i}: expected {old!r}, found {r}')
        sp = AT.sub(lambda m: f'<a:t>{escape(new)}</a:t>', sp, count=1)
        if color: sp = re.sub(r'srgbClr val="[0-9A-F]{6}"', f'srgbClr val="{color}"', sp)
        if '\n' in new:
            m = re.search(r'<a:r>(<a:rPr\b[^>]*/>|<a:rPr\b.*?</a:rPr>)<a:t>([^<]*)</a:t></a:r>', sp, re.S)
            if not m: raise SystemExit(f'{self.path} shape {i}: no run to break')
            sp = sp[:m.start()] + '<a:br/>'.join(f'<a:r>{m.group(1)}<a:t>{p}</a:t></a:r>' for p in m.group(2).split('\n')) + sp[m.end():]
        self._put(i, sp)
    def box(self, i):
        """A shape's (x, y, cx, cy) in EMU."""
        a, b, sp = self._shape(i); m = re.search(r'<a:off x="(\d+)" y="(\d+)"/><a:ext cx="(\d+)" cy="(\d+)"/>', sp)
        if not m: raise SystemExit(f'{self.path} shape {i}: no position')
        return tuple(int(v) for v in m.groups())
    def geo(self, i, **kw):
        """Move or resize shape i — only the x / y / cx / cy given change."""
        a, b, sp = self._shape(i); m = re.search(r'<a:off x="(\d+)" y="(\d+)"/><a:ext cx="(\d+)" cy="(\d+)"/>', sp)
        if not m: raise SystemExit(f'{self.path} shape {i}: no position')
        v = dict(zip(('x', 'y', 'cx', 'cy'), m.groups())); v.update({k: str(int(n)) for k, n in kw.items()})
        self._put(i, sp[:m.start()] + f'<a:off x="{v["x"]}" y="{v["y"]}"/><a:ext cx="{v["cx"]}" cy="{v["cy"]}"/>' + sp[m.end():])
    def sub(self, i, old, new, count=None):
        """Replace a literal inside shape i's XML (a font size, a paragraph spacing), refusing when it is absent — or,
        with `count`, when it does not occur exactly that many times."""
        a, b, sp = self._shape(i); n = sp.count(old)
        if n == 0 or (count is not None and n != count): raise SystemExit(f'{self.path} shape {i}: {old!r} occurs {n} times')
        self._put(i, sp.replace(old, new))
    def shift_items(self, lo, hi, first_y, old_pitch, new_pitch):
        """Re-space a column of repeated items (icon, heading, sub-line — shapes AND pictures): every element whose top is
        in [lo, hi) belongs to item k = (y − first_y) // old_pitch and moves by k × (new_pitch − old_pitch)."""
        def mv(m):
            y = int(m.group(2))
            if not lo <= y < hi: return m.group(0)
            k = (y - first_y) // old_pitch; return f'<a:off x="{m.group(1)}" y="{y + k * (new_pitch - old_pitch)}"/>'
        self.x = re.sub(r'<a:off x="(\d+)" y="(\d+)"/>', mv, self.x)
    def add_banner(self, text):
        """A closing banner of the deck's own style (the navy bar at y = 4,160,520), inserted before the footer so the
        slide ends the way every other content slide does. `text` is a plain sentence — the figures stay in the rows."""
        ids = [int(v) for v in re.findall(r'<p:cNvPr id="(\d+)"', self.x)]; i, j = max(ids) + 1, max(ids) + 2
        bar = (f'<p:sp><p:nvSpPr><p:cNvPr id="{i}" name="Banner"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="457200" y="4160520"/>'
               '<a:ext cx="8229600" cy="548640"/></a:xfrm><a:prstGeom prst="roundRect"><a:avLst><a:gd name="adj" fmla="val 16667"/></a:avLst></a:prstGeom>'
               f'<a:solidFill><a:srgbClr val="{NAVY}"/></a:solidFill><a:ln w="12700"><a:solidFill><a:srgbClr val="{NAVY}"/></a:solidFill><a:prstDash val="solid"/></a:ln></p:spPr></p:sp>'
               f'<p:sp><p:nvSpPr><p:cNvPr id="{j}" name="Banner text"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="640080" y="4160520"/>'
               '<a:ext cx="7863840" cy="548640"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr><p:txBody>'
               '<a:bodyPr wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" rtlCol="0" anchor="ctr"/><a:lstStyle/><a:p><a:pPr indent="0" marL="0"><a:buNone/></a:pPr>'
               '<a:r><a:rPr lang="en-US" sz="1400" b="1" dirty="0"><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill><a:latin typeface="Inter" pitchFamily="34" charset="0"/>'
               f'<a:ea typeface="Inter" pitchFamily="34" charset="-122"/><a:cs typeface="Inter" pitchFamily="34" charset="-120"/></a:rPr><a:t>{escape(text)}</a:t></a:r></a:p></p:txBody></p:sp>')
        foot = [(a, b) for a, b in spans(self.x) if '<a:off x="457200" y="4800600"/>' in self.x[a:b]]
        if len(foot) != 1 or '<a:off x="457200" y="4160520"/>' in self.x: raise SystemExit(f'{self.path}: no footer, or a banner already')
        a = foot[0][0]; self.x = self.x[:a] + bar + self.x[a:]
    def runs(self, i, old, new):
        """Replace every run of a shape in order (a table or a list with the same number of items)."""
        a, b, sp = self._shape(i); r = runs_of(sp)
        if [escape(o) for o in old] != r: raise SystemExit(f'{self.path} shape {i}: expected {old}, found {r}')
        it = iter(new); sp = AT.sub(lambda m: f'<a:t>{escape(next(it))}</a:t>', sp); self._put(i, sp)
    def paras(self, i, first, items):
        """Rebuild a list shape's paragraphs: each item is (template paragraph index, text). `first` checks the source."""
        a, b, sp = self._shape(i); ps = re.findall(r'<a:p>.*?</a:p>', sp, re.S)
        if runs_of(ps[0]) != [escape(first)]: raise SystemExit(f'{self.path} shape {i}: expected {first!r}')
        new = ''.join(AT.sub(lambda m, t=t: f'<a:t>{escape(t)}</a:t>', ps[k], count=1) for k, t in items)
        s, e = sp.index(ps[0]), sp.rindex(ps[-1]) + len(ps[-1]); self._put(i, sp[:s] + new + sp[e:])
    def save(self): open(self.path, 'w', encoding='utf-8').write(self.x)

def notes(d, n, old_start, new):
    """Replace a slide's speaker notes (the long run), checking how the old ones began."""
    p = f'{d}/ppt/notesSlides/notesSlide{n}.xml'; x = open(p, encoding='utf-8').read()
    rs = sorted(runs_of(x), key=len)
    if not rs[-1].startswith(escape(old_start)): raise SystemExit(f'{p}: notes begin {rs[-1][:60]!r}')
    x = x.replace(f'<a:t>{rs[-1]}</a:t>', f'<a:t>{escape(new)}</a:t>', 1); open(p, 'w', encoding='utf-8').write(x)

def rename_everywhere(d):
    """Last: the name and code wherever they remain (footers, headers, notes, document title). The managers' footer also
    takes the colleague deck's wording ("December 2026 link proposal"), so the two decks' footers read the same way."""
    for root, _, files in os.walk(d):
        for f in files:
            if not f.endswith('.xml'): continue
            p = os.path.join(root, f); x = open(p, encoding='utf-8').read()
            y = x.replace('Familiar Nine', NAME).replace('F9-24-K31s', CODE).replace('December 2026 CEA link · for managers', 'December 2026 link proposal · for managers')
            if y != x: open(p, 'w', encoding='utf-8').write(y)

# ── layout (4 Oct 2026, from a visual review of all 75 slides) ────────────────────────────────────────────
# Geometry the six shipped decks inherit from this template, so one change here reaches all of them; wording that
# depends on a design's own figures (tile captions, the Why lines) is deck-polish.py's.
ROW_Y0, ROW_CY, ROW_PITCH = 1801368, 365760, 438912                                   # one row: 0.40in tall, 0.08in apart
LABEL_X, LABEL_CX, TODAY_X, DESIGN_X, VALUE_CX = 640080, 4023360, 4663440, 6675120, 2011680
BANNER_Y = 4160520

def unify_rows(s):
    """Every table slide draws the SAME row — one height, one spacing, one set of columns — so nothing jitters when paging
    from a 3-row slide to a 5-row one. Five rows end 0.26in above the banner; three leave air above it, by design. Rows
    are found by their geometry (a full-width background, then its label, Today and design shapes), never by text."""
    headers = rows = pending = 0
    for i in range(len(spans(s.x))):
        x, y, cx, cy = s.box(i)
        if (y, cy) == (1417320, 320040):                                   # the Today / design column headings
            s.geo(i, x=TODAY_X if headers == 0 else DESIGN_X, cx=VALUE_CX); headers += 1
        elif (x, cx) == (457200, 8229600) and 1_700_000 < y < BANNER_Y:   # a row's background
            s.geo(i, y=ROW_Y0 + rows * ROW_PITCH, cy=ROW_CY); rows += 1; pending = 3
        elif pending:
            col = ((LABEL_X, LABEL_CX), (TODAY_X, VALUE_CX), (DESIGN_X, VALUE_CX))[3 - pending]
            s.geo(i, x=col[0], cx=col[1], y=ROW_Y0 + (rows - 1) * ROW_PITCH, cy=ROW_CY); pending -= 1
    if headers != 2 or not 3 <= rows <= 5 or pending: raise SystemExit(f'{s.path}: {headers} headings, {rows} rows, {pending} cells unplaced')

def widen_tile_text(s, inset):
    """The short-version tiles: the caption and sub-caption boxes take the tile's width less `inset` each side (they had
    0.1in more), so a deliberate two-word line fits without shrinking the type."""
    tiles = [i for i in range(len(spans(s.x))) if s.box(i)[1] == 1463040]
    if len(tiles) not in (4, 5): raise SystemExit(f'{s.path}: {len(tiles)} tiles')
    for i in tiles:
        x, _, cx, _ = s.box(i)
        for j in (i + 2, i + 3): s.geo(j, x=x + inset, cx=cx - 2 * inset)

def relayout_tiles(s):
    """The trade-offs slide: three rows of two tiles end 0.19in above a banner instead of running on towards the footer.
    A tile is 0.90in tall (was 1.07in) with 0.08in between; its title sits 0.08in down and its two-line body 0.38in."""
    OLD_Y, OLD_PITCH, NEW_CY, NEW_PITCH = 1371600, 1085000, 822960, 896112
    OFFS = {0: (0, NEW_CY), 158452: (73152, 274320), 469348: (347472, 347472)}   # old offset in the tile → (new offset, new height)
    for i in range(len(spans(s.x))):
        y = s.box(i)[1]
        if not OLD_Y <= y < OLD_Y + 3 * OLD_PITCH: continue
        k, off = divmod(y - OLD_Y, OLD_PITCH)
        if off not in OFFS: raise SystemExit(f'{s.path} shape {i}: unexpected tile offset {off}')
        s.geo(i, y=OLD_Y + k * NEW_PITCH + OFFS[off][0], cy=OFFS[off][1])

def restyle_table(s):
    """The At-a-glance table (a real <a:tbl>, the one on the slides): no vertical rules and a hairline between rows, so it
    reads like the row component above, and the label column indented to the rows' own left edge. The type stays 12pt:
    eleven rows in the space above the banner, and the closers row ("8h 10m · 8h 40m · 8h 10m") wraps at anything larger."""
    x, n = s.x, [0]
    for side in 'LR':
        x, c = re.subn(rf'<a:ln{side} w="9525".*?</a:ln{side}>', f'<a:ln{side}><a:noFill/></a:ln{side}>', x, flags=re.S); n[0] += c
    x, c = re.subn(r'<a:ln([TB]) w="9525"', r'<a:ln\1 w="6350"', x); n[0] += c
    cells = re.findall(r'<a:tc>.*?</a:tc>', x, re.S)
    if len(cells) % 3 or n[0] < 4 * len(cells): raise SystemExit(f'{s.path}: {len(cells)} cells, {n[0]} borders restyled')
    for i, c in enumerate(cells):
        if i % 3 == 0:
            if 'marL="91440"' not in c: raise SystemExit(f'{s.path}: label cell {i} has no left margin to widen')
            x = x.replace(c, c.replace('marL="91440"', f'marL="{LABEL_X - 457200}"'), 1)
    s.x = x

def bullet_font(d):
    """Every • bullet names its font: with none given, the renderer falls back to a symbol face and draws a small raised
    dot. Inter's own bullet sits centred on the x-height."""
    for n in os.listdir(f'{d}/ppt/slides'):
        if not n.endswith('.xml'): continue
        p = f'{d}/ppt/slides/{n}'; x = open(p, encoding='utf-8').read()
        y = x.replace('<a:buSzPct val="100000"/><a:buChar', '<a:buSzPct val="100000"/><a:buFont typeface="Inter" pitchFamily="34" charset="0"/><a:buChar')
        if y != x: open(p, 'w', encoding='utf-8').write(y)

def build(src, out, edit):
    d = tempfile.mkdtemp(); zipfile.ZipFile(src).extractall(d)
    edit(d); bullet_font(d); rename_everywhere(d)
    left = [f for r, _, fs in os.walk(d) for f in fs if f.endswith('.xml') and re.search(r'Familiar Nine|F9-24|24 weeks', open(os.path.join(r, f), encoding='utf-8').read())]
    if left: raise SystemExit(f'{out}: Familiar Nine text left in {left}')
    os.makedirs(os.path.dirname(out), exist_ok=True)
    if os.path.exists(out): os.remove(out)
    with zipfile.ZipFile(src) as zs, zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as z:
        for n in zs.namelist(): z.write(os.path.join(d, n), n)   # the source's own part order, [Content_Types].xml first
    shutil.rmtree(d)

# ── the colleague deck ────────────────────────────────────────────────────────────────────────────────────
def colleagues(d):
    S = lambda n: Slide(f'{d}/ppt/slides/slide{n}.xml')
    s = S(1); s.text(3, 'A proposal for discussion · CEA link · 24 weeks · code F9-24-K31s', f'A proposal for discussion · CEA link · 26 weeks · code {CODE}'); s.save()

    s = S(2)
    s.text(11, '6 in 24', '7 in 26'); s.text(15, '8 of 16', '9 of 15'); s.text(17, '8 new ones', '6 new ones')
    s.text(23, 'Same hours, shorter closers, more full weekends off — and half the times familiar.', 'Same hours, shorter closers, more full weekends off — and most of the times familiar.')
    s.save()
    notes(d, 2, 'This is the whole pitch', 'This is the whole pitch on one slide. Same 35-hour week as now. Never more than six days in a row. Seven full weekends off in every 26 weeks, against four in 20 today — though each person works about two more Saturdays a year. Nine of the fifteen shift times are ones people already work. And every closing shift is shorter: the weekday and Saturday closers by 30 minutes, the Sunday closer by 45.')

    s = S(3); s.text(7, 'It grows from 20 weeks to 24, so four more colleagues join it.', 'It grows from 20 weeks to 26, so six more colleagues join it.'); s.save()
    notes(d, 3, 'The change is coming', 'The change is coming whatever we do: the December 2026 timetable and the new staffing levels mean the link has to be redrawn. The question is which version we want. The staffing levels and the Sunday cover were confirmed verbally on 29 September 2026, and the 26-person link, its five cover weeks and a ceiling of 219 contracted days a year on 1 October 2026: nine rules the proposal is scored on. Fourteen on a Saturday is one of three flexible aims, worked towards but not scored.')

    s = S(4)
    s.text(11, '4.17', '4.19', NAVY); s.text(15, '217', '219', NAVY)
    s.text(17, 'No. Exactly the same 35-hour week — and about 2 fewer contracted days a year.', 'No. Exactly the same 35-hour week — and no more contracted days than today.')
    s.save()
    notes(d, 4, 'Monday to Saturday', 'Monday to Saturday, counting a cover week as four days worked, the same way the clerks do. Sundays are overtime by agreement, so they are not in these figures — today or in Second Nature. Exactly, Second Nature works out at 218.6 days a year: just under today’s 219, which is also the most the 26-week link is allowed (agreed 1 October 2026). The average shift is about two minutes longer, which is on slide 6.')

    s = S(5)
    s.text(11, '30', '28'); s.text(15, '22', '20')
    s.text(17, 'No. The same 35-hour week — plus about four more Saturdays at time and a quarter.', 'No. The same 35-hour week — plus about two more Saturdays at time and a quarter.')
    s.save()
    notes(d, 5, 'Basic pay follows', 'Basic pay follows the 35-hour week, which does not change; overtime and premium duties vary person to person. The Saturday and Sunday counts are the link’s fixed duties and leave out anything later given in a cover week. Saturdays in the contract are paid at time and a quarter, and on average a person is rostered about 28 a year instead of 26. Sundays stay overtime: 10 a week shared among 26 people instead of 8 among 20, so about the same each (20 a year, against 21). These are averages across the whole link — your own depends on which line you start on — and they are not a payslip; the Pay Calculator in the app does that.')

    s = S(6)
    s.text(1, 'Every closing shift gets shorter — the average goes up six minutes', 'Every closing shift gets shorter — the average goes up about two minutes')
    s.text(23, '8h 20m', '8h 16m', AMBER)
    s.save()
    notes(d, 6, 'Lead with the closers', 'Lead with the closers: they are the shifts people dislike most, and every one of them is shorter. Weekday 15:45 to 23:55 instead of 15:15 (8h 10m, not 8h 40m); Saturday 15:15 instead of 14:45 (8h 40m, not 9h 10m); Sunday 15:15 to 23:25 instead of 14:30 (8h 10m, not 8h 55m). Be upfront about the two minutes on the average: the 26-week link may have no more than 219 contracted days a year, so the same 35 hours go into slightly fewer, slightly longer shifts. Nothing runs over 9 hours.')

    s = S(7)
    s.text(1, 'Yes — about 59 a year each, against 39 today', 'A little — about 42 a year each, against 39 today')
    s.text(7, '59', '42', AMBER); s.text(11, '91', '94', AMBER); s.text(15, '65', '60', NAVY)
    s.text(21, 'More late finishes, only partly from the rules — but every closing shift is shorter.', 'Three more late finishes a year — the fewest the rules allow — and every closer is shorter.')
    s.save()
    notes(d, 7, 'Be straight about this one', 'Be straight about this one. About three more a year each finish at 23:00 or later, and all of that comes from the December rules: three on duty to the close every day, where today’s weekdays have two. In Second Nature the only people who finish after 23:00 are those three closers, so no link that meets the rules on 26 weeks can have fewer. Finishes at 22:00 or later go up more, from 78 to 94 a year: the rules also want five still on duty after 22:00, and two weekday lates finish at 22:30 with the ticket office. Early 06:20 starts hardly change. And because the shortest gap between two shifts is 14h 20m, nobody finishes late and starts early the next morning.')

    s = S(8)
    s.text(1, '16 shift times in all — 8 of them you already work', '15 shift times in all — 9 of them you already work')
    s.text(4, 'You work these today (8)', 'You work these today (9)')
    # the list is set at the right card's size (it was two points larger), and one time per line fills the card to match
    s.sub(5, 'sz="1200"', 'sz="1050"')
    s.paras(5, '06:20–13:35 · 06:20–14:20', [(0, t) for t in ('06:20–14:00', '06:20–14:20', '06:20–14:50', '07:15–15:45', '08:00–16:30', '12:00–20:00', '14:00–22:30', '14:30–22:00', '15:15–23:55')])
    s.text(8, 'New times (8)', 'New times (6)')
    s.paras(9, 'Weekdays', [(0, 'Weekdays'), (1, '07:00–16:00 — a new early, 9 hours'), (2, '15:00–22:30 — a new late, ending with the ticket office'),
                            (2, '15:45–23:55 — the closer, 30 minutes later than 15:15'), (3, 'Saturday'), (4, 'nothing new — every Saturday time is worked today'),
                            (6, 'Sunday'), (8, '15:15–23:25 — the closer, 45 minutes later than 14:30'), (9, '09:00–18:00 — a new early, 9 hours'),
                            (10, '14:30–22:30 — the ticket-office late')])
    s.text(11, 'Fewer shift times than today (16, not 18) — and most new ones are a small nudge.', 'Fewer shift times than today (15, not 18) — and nothing new on a Saturday.')
    s.save()
    notes(d, 8, 'Familiar Nine was built', f'Second Nature was built keeping the times people already work wherever that cost little. Nine of its fifteen are worked today, and it uses fewer different times overall than today’s eighteen. Several of the new times are shorter versions of today’s lates: the weekday and Sunday closers start later, and the Sunday office late finishes earlier. The full list is on the slide. The Sunday opener stays today’s 07:15–15:45.')

    s = S(9)
    s.text(7, '6 in 24', '7 in 26', GREEN); s.text(15, '22 of 24', '21 of 25', GREEN); s.text(23, '14 of 20', '15 of 21', GREEN)
    s.text(25, 'Six weekends off in 24 weeks, never more than six days on, and more rest.', 'Seven weekends off in 26 weeks, never more than six days on, and more rest.')
    s.save()
    notes(d, 9, 'A full weekend is', 'A full weekend is Saturday and the following Sunday both off. Today that is 4 weekends in a 20-week cycle; Second Nature gives 7 in 26. They fall unevenly, not every fourth week, and can be up to ten weeks apart, against seven today — worth saying, because somebody will start on the long gap. Most weeks now keep the same shift time Monday to Friday — 15 of the 21 working weeks — and all but two are all earlies or all lates, where today seven weeks in 16 mix the two. Rest days mostly come in pairs — 21 of the 25 rest-day breaks are two days or more, against 13 of 17 today — and there are no six-day weeks (today there is one).')

    s = S(10)
    s.text(1, 'The balance you have now stays the same', 'The balance you have now stays much the same')
    s.text(4, 'Half earlies, half lates', 'About half earlies, half lates')
    s.text(5, 'The same split as today: every early Monday to Saturday is matched by a late.', 'Saturday and Sunday split evenly; a weekday has one more late than early.')
    s.text(15, 'No shift in the balance of earlies and lates. Sundays stay overtime and still finish at 23:25.', 'Much the same balance. Sundays stay overtime and still finish at 23:25.')
    s.save()
    notes(d, 10, 'Earlies here means', 'Earlies here means a start before 11:00. A weekday has 7 earlies and 8 lates; a Saturday 7 and 7; a Sunday 5 and 5. Sundays: today 8 Sunday duties across 16 working weeks, Second Nature 10 across 21 — about the same share of working weeks. With six more people to share them, that is about one fewer Sunday a year each (20 against 21).')

    s = S(11)
    s.text(11, '23.6', '23.4', NAVY)
    s.save()
    notes(d, 11, 'Leave is only needed', 'Leave is only needed on working days Monday to Saturday; rest days and Sundays cost nothing. Today there is one unusual stretch of the roster where 14 days buys 30 days off — Second Nature spreads rest days more evenly, so its best is 28. On average it is the same, 23.4 days, and at its worst it is a day better than today. Four full weeks off (Sunday to Saturday) takes at least 15 days of leave, instead of 14 at best today. The exact averages depend on where you count a booking as starting, so quote them as “about the same”.')

    s = S(12)
    s.text(7, '14', '15', GREEN)
    s.save()

    s = S(13)
    s.text(3, 'Eight new shift times to learn', 'Six new shift times to learn')
    # every tile body runs to two lines (these two ran to one, beside four that did not), so the six tiles read as one set
    s.text(4, 'Most are within 10–30 minutes of a time we already work.', 'Several of them are shorter versions of the lates you already work today.')
    s.text(6, 'Shifts 6 minutes longer on average', 'Shifts about 2 minutes longer on average')
    s.text(7, 'That is what buys two fewer contracted days.', 'That is what keeps days at work to 219 a year or fewer on the same 35 hours.')
    s.text(10, 'Average and worst: about the same as today.', 'Average as today, and the worst a day better.')
    s.text(12, 'Four cover weeks remain', 'Five cover weeks')
    s.text(13, 'Evenly spaced — every sixth week (1, 7, 13, 19).', 'Spaced evenly: weeks 1, 6, 11, 16 and 21.')
    s.text(15, 'More late finishes', 'A few more late finishes')
    s.text(16, 'About 59 a year each at 23:00 or later on the fixed rota, against 39 today.', '42 a year each at 23:00 or later (today 39) — the rules’ minimum.')
    s.save()

    s = S(14)
    s.runs(2, ['Today', 'Familiar Nine', 'Contracted hours a week', '35h', '35h', 'Contracted days at work a year', '219', '217', 'Average shift · longest shift', '8h 14m · 9h 10m', '8h 20m · 9h 00m', 'Closing shift: weekday · Saturday · Sunday', '8h 40m · 9h 10m · 8h 55m', '8h 10m · 8h 40m · 8h 10m', 'Full weekends off', '4 in 20', '6 in 24', 'Most days worked in a row', '7 (up to 9)', '6', 'Shortest gap between two shifts', '12h 30m', '14h 20m', 'Different shift times', '18', '16 (8 you know)', 'On duty weekday · Saturday · Sunday', '11–12 · 10 · 8', '14 · 14 · 10', 'Avoidable fatigue warnings in the fixed rota', '4 (up to 5)', '0'],
               ['Today', NAME, 'Contracted hours a week', '35h', '35h', 'Contracted days at work a year', '219', '219', 'Average shift · longest shift', '8h 14m · 9h 10m', '8h 16m · 9h 00m', 'Closing shift: weekday · Saturday · Sunday', '8h 40m · 9h 10m · 8h 55m', '8h 10m · 8h 40m · 8h 10m', 'Full weekends off', '4 in 20', '7 in 26', 'Most days worked in a row', '7 (up to 9)', '6', 'Shortest gap between two shifts', '12h 30m', '14h 20m', 'Different shift times', '18', '15 (9 you know)', 'On duty weekday · Saturday · Sunday', '11–12 · 10 · 8', '15 · 14 · 10', 'Avoidable fatigue warnings in the fixed rota', '4 (up to 5)', '0'])
    s.text(4, 'Same hours, fewer contracted days, shorter closers, more full weekends off.', 'Same hours and days, shorter closers, more full weekends off.')
    restyle_table(s)
    s.save()

    # ── layout (4 Oct 2026): geometry only; the text edits above address shapes by index, so this runs after them ──
    for n in (4, 5, 6, 7, 9, 11, 12): s = S(n); unify_rows(s); s.save()
    s = S(2); widen_tile_text(s, 54864); s.save()
    # slides 3 and 13 ended 0.3in above the footer with no banner; their items close up and a banner (a plain sentence
    # from the speaker notes, no figure) ends them the way every other content slide ends
    s = S(3); s.shift_items(1_300_000, 4_600_000, 1335024, 862368, 713232)
    s.add_banner('The change is coming whatever we do — the question is which version we want.'); s.save()
    s = S(13); relayout_tiles(s); s.add_banner('None of these is hidden: each one is on the full proposal sheet too.'); s.save()

# ── the managers' deck ────────────────────────────────────────────────────────────────────────────────────
def managers(d):
    S = lambda n: Slide(f'{d}/ppt/slides/slide{n}.xml')
    s = S(1); s.text(3, 'For managers · 24 weeks · code F9-24-K31s · the full proposal sheet has the detail', f'For managers · 26 weeks · code {CODE} · the full proposal sheet has the detail'); s.save()

    s = S(2); s.text(15, '24', '26'); s.save()
    notes(d, 2, 'The nine rules', 'The nine rules are the December 2026 rules, as confirmed verbally on 29 September 2026: at least four at the open, three to the close, five after 22:00, 10 on a Sunday, the ticket office in fixed pairs, two on the floor at all times, handovers, 8–9 hour Sundays and no more shift times than today. Behind them sit three flexible rules that are aimed for when designing but not scored on: 14 on a Saturday, five cover weeks as evenly spaced as 26 weeks allow, and weekday closers starting at 15:45. Second Nature meets all three. It also stays under the ceiling of 219 contracted days a year agreed on 1 October 2026, at 218.6.')

    s = S(3); s.text(7, '14 · 14 · 10', '15 · 14 · 10', GREEN); s.text(15, '6 · 5 · 5', '7 · 5 · 5', GREEN); s.text(19, '3 · 4 · 3', '3 · 3 · 3', GREEN); s.save()
    notes(d, 3, 'Today’s link meets', 'Today’s link meets 4 of the 9 December rules; Second Nature meets all 9, and the three flexible ones as well (14 on a Saturday, evenly spaced cover weeks, 15:45 weekday closers). Someone finishing at exactly 22:00 is not counted as still on after 22:00. The floor figure leaves out the ticket office, except its second person at the quiet ends of the day. Three to the close every day is the minimum, and Second Nature rosters exactly that, which is what keeps its late finishes down.')

    s = Slide(f'{d}/ppt/slides/slide4.xml'); s.save()
    notes(d, 4, 'The figure compares', 'The figure compares, hour by hour, the share of the day’s floor staff on duty with the share of the day’s train service, each train weighted by its length. 0 would be a perfect match. It is a guide to shape, not a staffing requirement — the rules on the previous slide are the requirement. Against today’s link, Second Nature’s score is a little over half on weekdays, about a fifth on a Saturday and about two-fifths on a Sunday; lower is closer.')
    # the chart: its cached values, and the workbook PowerPoint opens to edit it
    c = f'{d}/ppt/charts/chart1.xml'; x = open(c, encoding='utf-8').read()
    for old, new in (('<c:v>27.7</c:v>', '<c:v>29.1</c:v>'), ('<c:v>15</c:v>', '<c:v>12.7</c:v>'), ('<c:v>26.4</c:v>', '<c:v>29</c:v>')):
        if x.count(old) != 1: raise SystemExit(f'chart: {old} found {x.count(old)} times')
        x = x.replace(old, new)
    open(c, 'w', encoding='utf-8').write(x)
    xp = f'{d}/ppt/embeddings/Microsoft_Excel_Worksheet1.xlsx'; wb = openpyxl.load_workbook(xp); ws = wb.active
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == 'Familiar Nine': cell.value = NAME; col = cell.column
    for r, v in ((2, 29.1), (3, 12.7), (4, 29.0)): ws.cell(row=r, column=col).value = v
    wb.save(xp)

    s = S(5)
    s.text(7, '24', '26'); s.text(15, '4 — 1 week in 6', '5 — 1 week in 5', NAVY)
    s.text(21, 'Four more people, 35 hours a week on average — Sundays still rely on overtime.', 'Six more people, 35 hours a week on average — Sundays still rely on overtime.')
    s.save()
    notes(d, 5, 'Working weeks average', 'Working weeks average exactly 35 hours Monday to Saturday — 44,100 minutes across the 21 working weeks — so no overtime is needed to balance the Monday-to-Saturday week. Individual weeks range from 25h 30m to 43h 40m; only the average is the contract, as today. Five cover weeks keep about the same share of the link as today’s four in 20 (one week in 5.2, against one in 5). Sunday is not contracted, so the ten Sunday duties are worked as rest-day overtime, exactly as today’s eight are. The two extra Sunday overtime duties a week come from the Sunday staffing of ten, confirmed verbally, not from Second Nature — any link that meets the rules needs them. The 26-week link, its five cover weeks and the ceiling of 219 contracted days a year were agreed on 1 October 2026; Second Nature gives 218.6.')

    s = S(7)
    s.text(1, 'It meets the brief while keeping 8 of its 16 shift times familiar', 'It meets the brief with the fewest late finishes the rules allow')
    s.text(7, 'Six full weekends off in 24 weeks — today four in 20', 'Seven full weekends off in 26 weeks — today four in 20')
    s.text(13, '8 of its 16 shift times are worked today', '9 of its 15 shift times are worked today')
    s.text(21, 'The trade-off: more late finishes — about 59 a year each, against 39 today.', 'The trade-off: three more late finishes a year each (42, today 39) — the rules’ minimum.')
    s.save()
    notes(d, 7, 'This is the case in one slide', f'This is the case in one slide. It meets every December staffing rule, and like {F9} it was built to keep the shift times people already work wherever that cost little — nine of its fifteen are worked today, and none on a Saturday is new. The honest cost is late finishes, and it is the smallest one possible: the December staffing needs three to the close every day, and in Second Nature only those three finish after 23:00, so finishes at 23:00 or later go from about 39 to 42 a year each. Every closing shift is 30 to 45 minutes shorter. Six shift times are new.')

    s = S(8)
    s.runs(5, ['Weekends off: 6 in 24 (today 4 in 20)', 'At most 6 days in a row (today 7)', 'Every closing shift 30–45 minutes shorter', '8 of the 16 shift times already worked', 'Shortest rest 14h 20m (today 12h 30m)'],
              ['Weekends off: 7 in 26 (today 4 in 20)', 'At most 6 days in a row (today 7)', 'Every closing shift 30–45 minutes shorter', '9 of the 15 shift times already worked', 'Shortest rest 14h 20m (today 12h 30m)'])
    s.runs(9, ['Late finishes: 59 a year (today 39)', '8 new shift times to learn', 'Average shift 6 minutes longer', 'Best 14-day leave stretch 28 days, not 30', 'Four weeks off takes 15 days’ leave at best, not 14'],
              ['Late finishes: 42 a year (today 39)', '6 new shift times to learn', 'Weekends off up to 10 weeks apart', 'Best 14-day leave stretch 28 days, not 30', 'Four weeks off needs 15 days’ leave, not 14'])   # one line: "…takes 15 days’ leave at best, not 14" left "not 14" alone on a second
    s.save()
    notes(d, 8, 'Days at work fall', 'Days at work stay at 219 a year (218.6 exactly; Monday to Saturday; Sunday is overtime), under the ceiling agreed on 1 October 2026. Late finishes average about three more a year each, all of them from the December staffing — three to the close every day, where today’s weekdays have two — so no link meeting the same rules on 26 weeks could have fewer. Full weekends off fall unevenly: seven in 26, but up to ten weeks apart. The average shift is about two minutes longer.')

    s = S(9)
    s.text(1, 'Three things confirmed verbally (29 Sep 2026), two decisions left', 'Three things confirmed verbally (29 Sep and 1 Oct 2026), two decisions left')
    s.text(8, '24 people on the link — confirmed', '26 people on the link — confirmed')
    s.text(9, 'Four more than today.', 'Six more than today, with five cover weeks.')
    s.save()

    # ── layout (4 Oct 2026): geometry only, after the text edits above (they address shapes by index) ──
    for n in (3, 5, 6): s = S(n); unify_rows(s); s.save()
    s = S(2); widen_tile_text(s, 68580); s.save()
    # the Why slide's six tick lines stopped 0.58in above the banner: 0.5in apart (was 0.44in) they stop 0.28in above it
    s = S(7); s.shift_items(1_300_000, 3_700_000, 1307592, 402336, 457200); s.save()
    # the two "notice" cards: each list takes the card's width less the same inset on both sides, and its five one-line
    # bullets sit 10pt apart (was 6pt), so neither card is a third empty
    s = S(8)
    for i in (5, 9): s.geo(i, cx=3657600); s.sub(i, '<a:spcPts val="600"/>', '<a:spcPts val="1000"/>', count=5)
    s.save()
    # Where it stands: five items close up from 0.7in to 0.6in apart and a banner (a plain sentence) ends the slide
    s = S(9); s.shift_items(1_280_000, 4_500_000, 1289304, 640080, 548640)
    s.add_banner('Colleagues’ views come first; the link is chosen after them, then who starts where.'); s.save()

if __name__ == '__main__':
    os.makedirs(OUT, exist_ok=True)
    build(f'{SRC}/Familiar-Nine-for-colleagues.pptx', f'{OUT}/Second-Nature-for-colleagues.pptx', colleagues)
    build(f'{SRC}/Familiar-Nine-for-managers.pptx', f'{OUT}/Second-Nature-for-managers.pptx', managers)
    print('wrote', OUT)
